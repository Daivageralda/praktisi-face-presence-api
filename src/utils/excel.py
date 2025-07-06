from pathlib import Path
from src.config import *
from openpyxl import Workbook, load_workbook

def update_Result(praktikan_id: str, result_label: str, file_path: Path = EXCEL_FILE):
    os.makedirs(EXCEL_FILE.parent, exist_ok=True)
    if not file_path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Presensi"
        ws.append(["Id Praktikan", "Presensi Berhasil", "Presensi Gagal", "Jumlah Presensi"])
    else:
        wb = load_workbook(file_path)
        ws = wb.active

    found = False
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == str(praktikan_id):
            if result_label == "Match":
                row[1].value = (row[1].value or 0) + 1
            else:
                row[2].value = (row[2].value or 0) + 1
            row[3].value = (row[3].value or 0) + 1
            found = True
            break

    if not found:
        ws.append([
            praktikan_id,
            1 if result_label == "Match" else 0,
            1 if result_label == "Not Match" else 0,
            1
        ])

    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    data_rows.sort(key=lambda x: int(x[0]))
    ws.delete_rows(2, ws.max_row)
    for row in data_rows:
        ws.append(row)

    wb.save(file_path)