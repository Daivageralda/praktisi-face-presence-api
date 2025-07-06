import tempfile
from pathlib import Path
from openpyxl import load_workbook
from src.utils import update_Result

def test_update_result_creates_new_file_and_adds_row():
    with tempfile.TemporaryDirectory() as tmpdir:
        file_path = Path(tmpdir) / "test_hasil.xlsx"

        # Pertama kali: file tidak ada, harus dibuat dan ditambahkan 1 baris
        update_Result("101", "Match", file_path)

        wb = load_workbook(file_path)
        ws = wb.active

        assert ws.title == "Presensi"
        assert ws.cell(row=2, column=1).value == "101"
        assert ws.cell(row=2, column=2).value == 1  # Match = berhasil
        assert ws.cell(row=2, column=3).value == 0
        assert ws.cell(row=2, column=4).value == 1

def test_update_result_updates_existing_id():
    with tempfile.TemporaryDirectory() as tmpdir:
        file_path = Path(tmpdir) / "test_hasil.xlsx"

        # Tambah pertama kali
        update_Result("102", "Not Match", file_path)
        # Tambah lagi dengan ID sama, tapi Match
        update_Result("102", "Match", file_path)

        wb = load_workbook(file_path)
        ws = wb.active

        assert ws.cell(row=2, column=1).value == "102"
        assert ws.cell(row=2, column=2).value == 1  # berhasil
        assert ws.cell(row=2, column=3).value == 1  # gagal
        assert ws.cell(row=2, column=4).value == 2  # jumlah

def test_update_result_adds_and_sorts_rows():
    with tempfile.TemporaryDirectory() as tmpdir:
        file_path = Path(tmpdir) / "test_hasil.xlsx"

        update_Result("5", "Match", file_path)
        update_Result("2", "Match", file_path)
        update_Result("3", "Not Match", file_path)

        wb = load_workbook(file_path)
        ws = wb.active

        # Cek urutan ID sudah ascending
        ids = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
        assert ids == sorted(ids, key=lambda x: int(x))
